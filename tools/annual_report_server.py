import argparse
import calendar
import hashlib
import io
import json
import os
import re
import tempfile
import threading
import zipfile
from datetime import date, timedelta
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse
from xml.etree import ElementTree as XML
from xml.sax.saxutils import escape

from arelle import Cntlr, ModelFormulaObject, ValidateXbrlCalcs
import reportlab
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


ROOT = Path(__file__).resolve().parent.parent
TAXONOMY = json.loads((ROOT / "data/annual-report/taxonomy-2026.json").read_text(encoding="utf-8"))
ARCHIVE = ROOT / "data/annual-report/et-gaap-2026.zip"
ROLE_CODES = ("101000", "201012", "301011")
XBRLI = "http://www.xbrl.org/2003/instance"
LINK = "http://www.xbrl.org/2003/linkbase"
XLINK = "http://www.w3.org/1999/xlink"
CORE_NAMESPACE = "http://xbrl.eesti.ee/taxonomy/et-gaap_2026-01-01/"
CORE_SCHEMA = "et-gaap-cor_2026-01-01.xsd"
VALIDATION_LOCK = threading.Lock()
FONT_DIRECTORY = Path(reportlab.__file__).parent / "fonts"
pdfmetrics.registerFont(TTFont("AnnualRegular", str(FONT_DIRECTORY / "Vera.ttf")))
pdfmetrics.registerFont(TTFont("AnnualBold", str(FONT_DIRECTORY / "VeraBd.ttf")))
pdfmetrics.registerFontFamily("AnnualRegular", normal="AnnualRegular", bold="AnnualBold")
with zipfile.ZipFile(ARCHIVE) as archive:
    LINKBASES = [name for name in archive.namelist()
                if any(name.endswith(f"_role-{code}.xml") for code in ROLE_CODES)
                or "/label/" in name and name.endswith(".xml")]


def form_rows(role):
    arcs = role["relationships"]
    children = {}
    for arc in arcs:
        children.setdefault(arc["parent"], []).append(arc)
    roots = set(children) - {arc["child"] for arc in arcs}
    rows = []
    visited = set()

    def visit(identifier, depth):
        if identifier in visited:
            return
        visited.add(identifier)
        concept = TAXONOMY["concepts"][identifier]
        if concept["abstract"] or concept["type"] == "xbrli:monetaryItemType":
            rows.append({**concept, "depth": depth, "label": concept["labels"].get("et") or concept["name"]})
        for arc in sorted(children.get(identifier, []), key=lambda item: item["order"]):
            visit(arc["child"], depth + 1)

    for identifier in sorted(roots):
        visit(identifier, 0)
    return rows


FORMS = [{"code": code, "title": role["definition"], "rows": form_rows(role)}
         for code in ROLE_CODES[1:]
         for role in TAXONOMY["roles"] if role["uri"].endswith(f"role-{code}")]
CONCEPTS = {row["id"]: row for form in FORMS for row in form["rows"] if not row["abstract"]}


def previous_year(value):
    try:
        return value.replace(year=value.year - 1)
    except ValueError:
        return value.replace(year=value.year - 1, day=28)


def check_payload(payload):
    if payload.get("profile") != "small-ou":
        raise ValueError("Unsupported report profile")
    if not re.fullmatch(r"\d{8}", str(payload.get("registryCode", ""))):
        raise ValueError("Sisestage 8-kohaline registrikood")
    if not str(payload.get("companyName", "")).strip():
        raise ValueError("Ettevotte nimi puudub")
    start = date.fromisoformat(payload["start"])
    end = date.fromisoformat(payload["end"])
    month_index = start.month + 17
    limit_year = start.year + month_index // 12
    limit_month = month_index % 12 + 1
    limit_day = min(start.day, calendar.monthrange(limit_year, limit_month)[1])
    maximum_end = date(limit_year, limit_month, limit_day) - timedelta(days=1)
    if start > end or end > maximum_end:
        raise ValueError("Aruande periood ei ole korrektne")
    values = {}
    for identifier, periods in payload.get("facts", {}).items():
        if identifier not in CONCEPTS or not isinstance(periods, dict):
            raise ValueError("Aruanne sisaldab tundmatut taksonoomia elementi")
        values[identifier] = {}
        for period in ("current", "prior"):
            value = periods.get(period)
            if value in (None, ""):
                continue
            try:
                amount = Decimal(str(value))
            except InvalidOperation as error:
                raise ValueError("Summa ei ole korrektne") from error
            if not amount.is_finite() or abs(amount) > Decimal("1e15"):
                raise ValueError("Summa ei ole korrektne")
            values[identifier][period] = amount.quantize(Decimal("0.01"))
    for period in ("current", "prior"):
        for name in ("Assets", "LiabilitiesAndEquity", "TotalAnnualPeriodProfitLoss"):
            identifier = next(key for key, concept in CONCEPTS.items() if concept["name"] == name)
            if period not in values.get(identifier, {}):
                raise ValueError(f"Puudub {CONCEPTS[identifier]['label']} ({period})")
        assets = values["et-gaap_Assets"][period]
        liabilities = values["et-gaap_LiabilitiesAndEquity"][period]
        if assets != liabilities:
            raise ValueError(f"Bilanss ei ole tasakaalus ({period})")
    return start, end, values


def instance_xml(payload, local=False):
    start, end, values = check_payload(payload)
    for prefix, namespace in (("xbrli", XBRLI), ("link", LINK), ("xlink", XLINK), ("et-gaap", CORE_NAMESPACE)):
        XML.register_namespace(prefix, namespace)
    root = XML.Element(f"{{{XBRLI}}}xbrl", {"xmlns:iso4217": "http://www.xbrl.org/2003/iso4217"})
    schemas = [CORE_SCHEMA] + list(dict.fromkeys(role["schema"] for role in TAXONOMY["roles"]
                                               if any(role["uri"].endswith(f"role-{code}") for code in ROLE_CODES)))
    for schema_file in schemas:
        schema = f"{ARCHIVE}/{schema_file}" if local else f"http://xbrl.eesti.ee/taxonomy/{schema_file}"
        XML.SubElement(root, f"{{{LINK}}}schemaRef", {f"{{{XLINK}}}type": "simple", f"{{{XLINK}}}href": schema})
    for linkbase_file in LINKBASES:
        kind = "label" if "/label/" in linkbase_file else {"cal": "calculation", "def": "definition", "pre": "presentation"}[Path(linkbase_file).name[:3]]
        location = f"{ARCHIVE}/{linkbase_file}" if local else f"http://xbrl.eesti.ee/taxonomy/{linkbase_file}"
        XML.SubElement(root, f"{{{LINK}}}linkbaseRef", {f"{{{XLINK}}}type": "simple", f"{{{XLINK}}}href": location,
                                                      f"{{{XLINK}}}arcrole": "http://www.w3.org/1999/xlink/properties/linkbase",
                                                      f"{{{XLINK}}}role": f"http://www.xbrl.org/2003/role/{kind}LinkbaseRef"})
    for period, first, last in (("current", start, end), ("prior", previous_year(start), previous_year(end))):
        for kind in ("instant", "duration"):
            context = XML.SubElement(root, f"{{{XBRLI}}}context", {"id": f"{period}-{kind}"})
            entity = XML.SubElement(context, f"{{{XBRLI}}}entity")
            XML.SubElement(entity, f"{{{XBRLI}}}identifier", {"scheme": "http://www.rik.ee"}).text = payload["registryCode"]
            interval = XML.SubElement(context, f"{{{XBRLI}}}period")
            if kind == "instant":
                XML.SubElement(interval, f"{{{XBRLI}}}instant").text = last.isoformat()
            else:
                XML.SubElement(interval, f"{{{XBRLI}}}startDate").text = first.isoformat()
                XML.SubElement(interval, f"{{{XBRLI}}}endDate").text = last.isoformat()
    unit = XML.SubElement(root, f"{{{XBRLI}}}unit", {"id": "EUR"})
    XML.SubElement(unit, f"{{{XBRLI}}}measure").text = "iso4217:EUR"
    metadata = {"CompanyName": payload["companyName"], "RegistryCode": payload["registryCode"],
                "AnnualReportName": f"Majandusaasta aruanne {end.year}", "BeginningAndEndOfAnnualPeriod": f"{start.isoformat()} - {end.isoformat()}"}
    for name, value in metadata.items():
        concept = next(concept for concept in TAXONOMY["concepts"].values() if concept["name"] == name)
        XML.SubElement(root, f"{{{concept['namespace']}}}{name}", {"contextRef": "current-instant"}).text = value
    for identifier, periods in values.items():
        concept = CONCEPTS[identifier]
        for period, amount in periods.items():
            XML.SubElement(root, f"{{{concept['namespace']}}}{concept['name']}",
                           {"contextRef": f"{period}-{concept['periodType']}", "unitRef": "EUR", "decimals": "2"}).text = str(amount)
    return XML.tostring(root, encoding="utf-8", xml_declaration=True)


def validate_instance(payload):
    instance = instance_xml(payload, local=True)
    with VALIDATION_LOCK, tempfile.TemporaryDirectory(prefix="arvesemu-xbrl-") as directory:
        path = Path(directory) / "annual-report.xbrl"
        path.write_bytes(instance)
        controller = Cntlr.Cntlr(logFileName="logToBuffer")
        controller.webCache.timeout = 15
        controller.modelManager.formulaOptions = ModelFormulaObject.FormulaOptions()
        controller.modelManager.validateCalcs = ValidateXbrlCalcs.ValidateCalcsMode.XBRL_v2_1
        try:
            model = controller.modelManager.load(str(path))
            controller.modelManager.validate()
            errors = [str(error) for error in model.errors]
            logs = json.loads(controller.logHandler.getJson()).get("log", [])
            blocking = [entry for entry in logs if str(entry.get("level", "")).lower() in ("error", "critical", "inconsistency")]
            return {"valid": not errors and not blocking, "errors": errors,
                    "messages": [entry.get("message", {}).get("text", "") for entry in logs if str(entry.get("level", "")).lower() in ("error", "warning", "critical", "inconsistency")]}
        finally:
            controller.close()


def workbook_bytes(payload):
    start, end, values = check_payload(payload)
    book = Workbook()
    info = book.active
    info.title = "Taksonoomia"
    for row in [("Ettevote", payload["companyName"]), ("Registrikood", payload["registryCode"]),
                ("Periood", f"{start} - {end}"), ("Taksonoomia", TAXONOMY["version"]), ("Allikas", TAXONOMY["archiveUrl"]),
                ("Staatus", "Mustand; lisad ja tegevusaruanne ei ole lisatud")]:
        info.append(row)
    for row in info:
        for cell in row:
            if isinstance(cell.value, str):
                cell.data_type = "s"
    for form in FORMS:
        sheet = book.create_sheet(form["code"])
        sheet.append([form["title"], end.isoformat(), previous_year(end).isoformat(), "XBRL element"])
        for concept in form["rows"]:
            periods = values.get(concept["id"], {})
            sheet.append([concept["label"], float(periods["current"]) if "current" in periods else None,
                          float(periods["prior"]) if "prior" in periods else None, concept["name"]])
            for cell in sheet[sheet.max_row]:
                if cell.column in (2, 3):
                    cell.number_format = "#,##0.00"
                if concept["abstract"]:
                    cell.font = Font(bold=True, color="234337")
        sheet.freeze_panes = "B2"
        for column, width in (("A", 65), ("B", 20), ("C", 20), ("D", 50)):
            sheet.column_dimensions[column].width = width
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="234337")
            cell.font = Font(color="FFFFFF", bold=True)
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()


def pdf_bytes(payload):
    start, end, values = check_payload(payload)
    stream = io.BytesIO()
    document = SimpleDocTemplate(stream, pagesize=A4, title=f"Majandusaasta aruanne {end.year}",
                                 author=payload["companyName"], leftMargin=18 * mm, rightMargin=18 * mm,
                                 topMargin=20 * mm, bottomMargin=22 * mm)
    normal = ParagraphStyle("AnnualBody", fontName="AnnualRegular", fontSize=9, leading=13, spaceAfter=8)
    heading = ParagraphStyle("AnnualHeading", parent=normal, fontName="AnnualBold", fontSize=15, leading=20, spaceAfter=12, keepWithNext=True)
    title = ParagraphStyle("AnnualTitle", parent=heading, fontSize=24, leading=30, spaceAfter=20)
    label = ParagraphStyle("AnnualLabel", parent=normal, fontSize=8, leading=11, spaceAfter=0)
    group = ParagraphStyle("AnnualGroup", parent=label, fontName="AnnualBold")
    format_date = lambda value: value.strftime("%d.%m.%Y")
    period_label = f"{format_date(start)} - {format_date(end)}"
    prior_label = f"{format_date(previous_year(start))} - {format_date(previous_year(end))}"
    money = lambda value: Decimal(str(value or 0)).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
    format_amount = lambda value: f"{money(value):,}".replace(",", " ")
    concept_id = lambda name: next((identifier for identifier, concept in CONCEPTS.items() if concept["name"] == name), "")
    report_value = lambda name, period="current": values.get(concept_id(name), {}).get(period, Decimal("0.00"))
    raw_address = str(payload.get("address", "")).strip()
    postal_match = re.search(r"(?<!\d)(\d{5})(?!\d)", raw_address)
    postal_code = postal_match.group(1) if postal_match else ""
    address_parts = [part.strip().strip(",") for part in re.split(r"[\n,]+", raw_address) if part.strip()]
    address_parts = [re.sub(r"\b\d{5}\b", "", part).strip(" ,") for part in address_parts]
    address_parts = [part for part in address_parts if part]
    county = next((part for part in address_parts if "maakond" in part.lower()), "")
    city = next((part for part in address_parts if "tallinn" in part.lower() or "tartu" in part.lower()), "")
    district = next((part for part in address_parts if "linnaosa" in part.lower()), "")
    other_address = [part for part in address_parts if part not in {county, city, district}]
    ordered_address = [part for part in (county, city, district, *other_address) if part]
    post_address = ", ".join(ordered_address)
    title_page = [Spacer(1, 16 * mm), Paragraph("MAJANDUSAASTA ARUANNE", title),
                  Paragraph(f"aruandeaasta algus: {format_date(start)}", normal),
                  Paragraph(f"aruandeaasta l\u00f5pp: {format_date(end)}", normal),
                  Spacer(1, 8 * mm), Paragraph(f"\u00e4rinimi: {escape(payload['companyName'])}", normal),
                  Paragraph(f"registrikood: {escape(payload['registryCode'])}", normal)]
    if post_address:
        title_page.append(Paragraph(f"postiaadress: {escape(post_address)}", normal))
    if postal_code:
        title_page.append(Paragraph(f"postisihtnumber: {escape(postal_code)}", normal))
    for key, label_text in (("phone", "telefon"), ("email", "e-posti aadress")):
        value = str(payload.get(key, "")).strip()
        if value:
            title_page.append(Paragraph(f"{label_text}: {escape(value)}", normal))
    title_page.extend([Spacer(1, 10 * mm), Paragraph("V\u00c4LJAV\u00d5TE · ALLKIRJASTAMATA", heading),
                       Paragraph("Kohaliku rakenduse koostatud eelvaade. Riigile esitatav l\u00f5plik PDF genereeritakse e-\u00c4riregistris p\u00e4rast vajalike vormide t\u00e4itmist ja allkirjastamist.", normal)])
    story = title_page + [PageBreak(), Paragraph(escape(payload["companyName"]), normal),
                          Paragraph(f"{end.year}. a. majandusaasta aruanne", normal),
                          Spacer(1, 8 * mm), Paragraph("Sisukord", heading)]
    contents = [("Raamatupidamise aastaaruanne", 3), ("Bilanss", 3), ("Kasumiaruanne", 4),
                ("Raamatupidamise aastaaruande lisad", 5), ("Aruande allkirjad", 6),
                ("Kasumi jaotamise ettepanek", 7), ("M\u00fc\u00fcgitulu jaotus tegevusalade l\u00f5ikes", 8)]
    story.extend(Paragraph(f"{escape(label_text)} <font color='#66736b'>{page_number}</font>", normal) for label_text, page_number in contents)

    def statement_table(form):
        current_caption = format_date(end) if form["code"] == "201012" else str(end.year)
        prior_caption = format_date(previous_year(end)) if form["code"] == "201012" else str(previous_year(end).year)
        rows = [[Paragraph("", group), Paragraph(current_caption, group), Paragraph(prior_caption, group)]]
        active = {identifier for identifier, periods in values.items() if any(amount != 0 for amount in periods.values())}
        pending_groups = []
        emitted_groups = set()
        group_rows = []
        for concept in form["rows"]:
            if concept["abstract"]:
                pending_groups = [item for item in pending_groups if item["depth"] < concept["depth"]] + [concept]
                continue
            if concept["id"] not in active:
                continue
            for parent in pending_groups:
                if parent["id"] not in emitted_groups:
                    group_rows.append(len(rows))
                    rows.append([Paragraph(escape(parent["label"]), group), "", ""])
                    emitted_groups.add(parent["id"])
            periods = values.get(concept["id"], {})
            rows.append([Paragraph(escape(concept["label"]), label), format_amount(periods.get("current")), format_amount(periods.get("prior"))])
        if len(rows) == 1:
            rows.append([Paragraph("Esitatavad kirjed puuduvad.", label), "", ""])
        table = Table(rows, colWidths=[document.width - 170, 85, 85], repeatRows=1)
        commands = [("FONTNAME", (0, 0), (-1, -1), "AnnualRegular"), ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                    ("GRID", (0, 0), (-1, -1), .45, colors.HexColor("#c9c9c9")),
                    ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ALIGN", (1, 1), (-1, -1), "RIGHT")]
        for row_index in group_rows:
            commands.extend([("BACKGROUND", (0, row_index), (-1, row_index), colors.HexColor("#f4f4f4")),
                             ("SPAN", (0, row_index), (-1, row_index))])
        table.setStyle(TableStyle(commands))
        return table

    balance_form, income_form = FORMS
    story.extend([PageBreak(), Paragraph("Raamatupidamise aastaaruanne", heading),
                  Paragraph("Bilanss", heading), Paragraph("(eurodes)", normal), statement_table(balance_form),
                  PageBreak(), Paragraph("Kasumiaruanne", heading), Paragraph("(eurodes)", normal), statement_table(income_form),
                  PageBreak(), Paragraph("Raamatupidamise aastaaruande lisad", heading),
                  Paragraph("Lisa 1 Arvestusp\u00f5him\u00f5tted", ParagraphStyle("AnnualNoteHeading", parent=normal, fontName="AnnualBold", fontSize=11, spaceAfter=8)),
                  Paragraph("K\u00e4esolev aruanne on Eesti finantsaruandluse standardist l\u00e4htuv mikroettev\u00f5tja l\u00fchendatud raamatupidamise aastaaruanne, mille eesm\u00e4rk on anda aruande kasutajale raamatupidamise seaduses n\u00f5utud informatsiooni aruandekohustuslase finantsseisundi ja -tulemuse kohta.", normal),
                  Paragraph("Ettev\u00f5tte tegevuse j\u00e4tkumise hinnang: esitamata.", normal),
                  PageBreak(), Paragraph("Aruande digitaalallkirjad", heading),
                  Paragraph("Aruanne on allkirjastamata. Digitaalallkirjad lisatakse e-\u00c4riregistris.", normal),
                  Paragraph("Osanike koosoleku kinnitamise staatus: kinnitamata.", normal),
                  PageBreak(), Paragraph("Kasumi jaotamise ettepanek", heading), Paragraph("(eurodes)", normal)])
    retained = report_value("RetainedEarningsLoss")
    current_profit = report_value("TotalAnnualPeriodProfitLoss")
    available_profit = retained + current_profit
    proposal_rows = [["N\u00e4itaja", format_date(end)],
                     ["Eelmiste perioodide jaotamata kasum (kahjum)", format_amount(retained)],
                     ["Aruandeaasta kasum (kahjum)", format_amount(current_profit)],
                     ["Kokku", format_amount(available_profit)],
                     ["Kasumi jaotamise ettepanek", "Esitamata"]]
    proposal_table = Table(proposal_rows, colWidths=[document.width - 120, 120])
    proposal_table.setStyle(TableStyle([("FONTNAME", (0, 0), (-1, -1), "AnnualRegular"), ("FONTSIZE", (0, 0), (-1, -1), 8),
                                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#eeeeee")),
                                        ("GRID", (0, 0), (-1, -1), .45, colors.HexColor("#c9c9c9")),
                                        ("ALIGN", (1, 1), (-1, -1), "RIGHT"), ("TOPPADDING", (0, 0), (-1, -1), 6),
                                        ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
    story.extend([proposal_table, Spacer(1, 5 * mm),
                  Paragraph("Jaotamise otsust ega dividendide jaotust ei tuletata raamatupidamiskannetest; see osa vajab juhatuse/osanike sisendit e-\u00c4riregistris.", normal),
                  PageBreak(), Paragraph("M\u00fc\u00fcgitulu jaotus tegevusalade l\u00f5ikes", heading),
                  Paragraph("Tegevusala", normal), Paragraph(f"M\u00fc\u00fcgitulu kokku: {format_amount(report_value('Revenue'))} EUR", normal),
                  Paragraph("EMTAK koodi ja tegevusalap\u00f5hist m\u00fc\u00fcgitulu jaotust ei ole Kontoplaanis seadistatud. RIK-i juhendi kohaselt tuleb esitada kuni 10 suuremat tegevusala v\u00f5i jaotada v\u00e4hemalt 90% m\u00fc\u00fcgitulust EMTAK koodide l\u00f5ikes.", normal)])

    def cover_page(canvas, current_document):
        canvas.saveState()
        canvas.setFont("AnnualRegular", 8)
        canvas.setFillColor(colors.HexColor("#66736b"))
        canvas.drawString(current_document.leftMargin, 12 * mm, "Eelvaade · allkirjastamata · mitte riigile esitamiseks")
        canvas.restoreState()

    def page_header_footer(canvas, current_document):
        canvas.saveState()
        canvas.setFont("AnnualRegular", 8)
        canvas.setFillColor(colors.HexColor("#222222"))
        canvas.drawString(current_document.leftMargin, A4[1] - 12 * mm, payload["companyName"])
        canvas.drawRightString(A4[0] - current_document.rightMargin, A4[1] - 12 * mm, f"{end.year}. a. majandusaasta aruanne")
        canvas.setFillColor(colors.HexColor("#66736b"))
        canvas.drawRightString(A4[0] - current_document.rightMargin, 12 * mm, f"{current_document.page}")
        canvas.restoreState()

    document.build(story, onFirstPage=cover_page, onLaterPages=page_header_footer)
    return stream.getvalue()


class Handler(BaseHTTPRequestHandler):
    def origin_allowed(self):
        origin = self.headers.get("Origin")
        allowed = set(os.environ.get("ANNUAL_REPORT_ALLOWED_ORIGINS", "").split())
        return not origin or urlparse(origin).hostname in ("localhost", "127.0.0.1") or origin in allowed

    def end_headers(self):
        origin = self.headers.get("Origin")
        if origin and self.origin_allowed():
            self.send_header("Access-Control-Allow-Origin", origin)
            self.send_header("Vary", "Origin")
        super().end_headers()

    def do_OPTIONS(self):
        if not self.origin_allowed():
            self.send_error(403)
            return
        self.send_response(204)
        self.send_header("Access-Control-Allow-Methods", "POST, GET, OPTIONS")
        self.send_header("Access-Control-Allow-Headers", "Content-Type")
        self.end_headers()

    def send_json(self, data, status=200):
        content = json.dumps(data, ensure_ascii=False).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(content)))
        self.end_headers()
        self.wfile.write(content)

    def do_GET(self):
        if self.path == "/api/annual-report/profile":
            self.send_json({"version": TAXONOMY["version"], "source": TAXONOMY["source"], "profile": "small-ou", "forms": FORMS})
        else:
            self.send_error(404)

    def do_POST(self):
        if not self.origin_allowed() or self.path not in ("/api/annual-report/check", "/api/annual-report/export"):
            self.send_error(403)
            return
        try:
            length = int(self.headers.get("Content-Length", "0"))
            if length <= 0 or length > 1000000:
                raise ValueError("Request size is invalid")
            payload = json.loads(self.rfile.read(length))
            check_payload(payload)
            if self.path.endswith("/check"):
                self.send_json(validate_instance(payload))
                return
            format_name = payload.get("format")
            if format_name == "xbrl":
                result = validate_instance(payload)
                if not result["valid"]:
                    self.send_json(result, 422)
                    return
                content, media = instance_xml(payload), "application/xml"
            elif format_name == "xlsx":
                content, media = workbook_bytes(payload), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            elif format_name == "pdf":
                content, media = pdf_bytes(payload), "application/pdf"
            else:
                raise ValueError("Unsupported format")
            self.send_response(200)
            self.send_header("Content-Type", media)
            self.send_header("Content-Disposition", f'attachment; filename="majandusaasta-aruanne-{payload["registryCode"]}-{payload["end"][:4]}.{format_name}"')
            self.send_header("Content-Length", str(len(content)))
            self.end_headers()
            self.wfile.write(content)
        except (ValueError, KeyError, TypeError) as error:
            self.send_json({"valid": False, "errors": [str(error)]}, 422)
        except Exception:
            self.send_json({"valid": False, "errors": ["Aruande genereerimine ebaonnestus"]}, 500)

    def log_message(self, format, *args):
        return


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--port", type=int, default=8765)
    parser.add_argument("--write-profile", type=Path)
    arguments = parser.parse_args()
    if hashlib.sha256(ARCHIVE.read_bytes()).hexdigest() != TAXONOMY["sha256"]:
        raise ValueError("Taxonomy archive checksum does not match the catalogue")
    if arguments.write_profile:
        arguments.write_profile.parent.mkdir(parents=True, exist_ok=True)
        arguments.write_profile.write_text(json.dumps({"version": TAXONOMY["version"], "source": TAXONOMY["source"],
                                                       "profile": "small-ou", "forms": FORMS}, ensure_ascii=False), encoding="utf-8")
        return
    server = ThreadingHTTPServer(("127.0.0.1", arguments.port), Handler)
    print(f"Arvesemu annual reports: http://127.0.0.1:{arguments.port}", flush=True)
    server.serve_forever()


if __name__ == "__main__":
    main()